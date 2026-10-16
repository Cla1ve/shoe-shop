"""Готовит учебные файлы для импорта: Products_import.xlsx, Stock_Items_import.xlsx и другие.

На экзамене такие файлы выдаёт заказчик, а у меня их не было, поэтому я собрал похожие сам.
Запуск из корня проекта:  python tools/make_import_files.py
Нужен openpyxl (только для этого скрипта, сам проект работает на встроенных модулях).
"""
import random
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font

random.seed(42)

# (наименование, категория, производитель, описание, состав, цена)
PRODUCTS = [
    ("Туфли-лодочки Classic", "Женская обувь", "Salamander", "Классические туфли на устойчивом каблуке 6 см", "Натуральная кожа, подкладка из кожи, подошва каучук", 8990),
    ("Ботильоны Chelsea Woman", "Женская обувь", "Rieker", "Ботильоны с эластичными вставками по бокам", "Натуральная кожа, текстильная подкладка, подошва ТЭП", 11990),
    ("Балетки Soft Flat", "Женская обувь", "Ecco", "Мягкие балетки на каждый день", "Мягкая натуральная кожа, кожаная стелька", 7490),
    ("Сапоги зимние Nordic", "Женская обувь", "Rieker", "Тёплые сапоги до колена", "Натуральная кожа, мех, подошва ТЭП", 14990),
    ("Босоножки Summer Strap", "Женская обувь", "Camper", "Летние босоножки с ремешком на щиколотке", "Натуральная кожа, подошва ЭВА", 9490),
    ("Лоферы City", "Женская обувь", "Ecco", "Лоферы для офиса и города", "Натуральная кожа, подошва полиуретан", 10490),
    ("Ботинки Trek Lady", "Женская обувь", "Salamander", "Демисезонные ботинки с мембраной", "Нубук, мембрана, подошва резина", 12990),
    ("Сандалии Beach", "Женская обувь", "Zenden", "Лёгкие сандалии для лета", "Искусственная кожа, текстиль", 3490),
    ("Туфли на каблуке Evening", "Женская обувь", "Zenden", "Вечерние туфли, каблук 8 см", "Искусственная кожа, подошва ТПУ", 4990),
    ("Кеды Canvas Woman", "Женская обувь", "Converse", "Классические кеды из ткани", "Текстиль, резиновая подошва", 5490),
    ("Ботинки Oxford Classic", "Мужская обувь", "Ecco", "Классические оксфорды со шнуровкой", "Натуральная кожа, кожаная подкладка", 15990),
    ("Ботинки Chelsea Man", "Мужская обувь", "Rieker", "Челси на каждый день", "Натуральная кожа, подошва ТЭП", 12490),
    ("Кроссовки Urban Walk", "Мужская обувь", "Skechers", "Повседневные кроссовки с мягкой стелькой", "Текстиль, амортизирующая стелька", 7990),
    ("Туфли Derby", "Мужская обувь", "Salamander", "Дерби с открытой шнуровкой", "Натуральная кожа", 13490),
    ("Ботинки зимние Arctic", "Мужская обувь", "Camper", "Зимние ботинки для сильных морозов", "Кожа, натуральный мех, подошва резина", 17990),
    ("Мокасины Drive", "Мужская обувь", "Ecco", "Мокасины для вождения", "Кожа, подошва резина", 10990),
    ("Сандалии Trek Sandal", "Мужская обувь", "Rieker", "Закрытые сандалии для походов", "Кожа, текстиль", 6490),
    ("Кеды Classic Low", "Мужская обувь", "Converse", "Низкие кеды", "Текстиль", 4990),
    ("Ботинки Work Pro", "Мужская обувь", "Zenden", "Рабочие ботинки с защитным подноском", "Кожа, стальной подносок", 9990),
    ("Ботинки Kids Winter", "Детская обувь", "Котофей", "Зимние ботинки для детей", "Мембрана, искусственный мех", 5990),
    ("Кроссовки Kids Run", "Детская обувь", "Adidas", "Лёгкие кроссовки для прогулок и спорта", "Сетка, синтетика", 4990),
    ("Сандалии Kids Summer", "Детская обувь", "Котофей", "Летние сандалии на липучке", "Искусственная кожа", 2990),
    ("Сапоги резиновые Rain", "Детская обувь", "Reima", "Резиновые сапоги для дождя", "Резина, текстильная подкладка", 2490),
    ("Ботинки School", "Детская обувь", "Ecco", "Школьные ботинки", "Натуральная кожа", 6990),
    ("Кеды Kids Canvas", "Детская обувь", "Converse", "Детские кеды", "Текстиль", 3490),
    ("Кроссовки Air Max 270", "Спортивная обувь", "Nike", "Кроссовки с воздушной подушкой Air", "Текстиль, синтетика, амортизация Air", 13990),
    ("Кроссовки Ultraboost Light", "Спортивная обувь", "Adidas", "Беговые кроссовки с подошвой Boost", "Primeknit, подошва Boost", 17990),
    ("Кроссовки Run Fast 3", "Спортивная обувь", "Puma", "Кроссовки для бега", "Сетка, пена ProFoam", 7490),
    ("Бутсы Predator Club", "Спортивная обувь", "Adidas", "Футбольные бутсы для искусственного газона", "Синтетика, шипы ТПУ", 6990),
    ("Кроссовки Pegasus 41", "Спортивная обувь", "Nike", "Универсальные беговые кроссовки", "Сетка, пена React", 12990),
    ("Кроссовки баскетбольные Court", "Спортивная обувь", "Nike", "Баскетбольные кроссовки с высоким верхом", "Кожа, текстиль", 11490),
    ("Кроссовки Trail Master", "Спортивная обувь", "Salomon", "Трейловые кроссовки с мембраной", "Мембрана, подошва Contagrip", 14490),
    ("Кроссовки Gel Rider", "Спортивная обувь", "Asics", "Кроссовки с гелевой амортизацией", "Сетка, гель", 9490),
]
SIZES = list(range(28, 47))
SIZE_RANGE = {"Женская обувь": (36, 41), "Мужская обувь": (40, 46), "Детская обувь": (28, 35), "Спортивная обувь": (36, 46)}

USERS = [
    ("Соколова Мария Андреевна", "masokolova", "Администратор"),
    ("Петров Алексей Николаевич", "anpetrov", "Менеджер"),
    ("Кузнецова Ольга Викторовна", "ovkuznetsova", "Менеджер"),
    ("Иванов Сергей Игоревич", "isivanov", "Клиент"),
    ("Смирнова Екатерина Дмитриевна", "edsmirnova", "Клиент"),
    ("Волков Дмитрий Олегович", "dovolkov", "Клиент"),
    ("Морозова Анна Павловна", "apmorozova", "Клиент"),
    ("Новиков Артём Сергеевич", "asnovikov", "Клиент"),
    ("Фёдорова Елена Ивановна", "eifedorova", "Клиент"),
    ("Лебедев Павел Андреевич", "palebedev", "Клиент"),
]
GUESTS = ["Громов Игорь Станиславович", "Орлова Татьяна Юрьевна"]


def header(ws, names):
    ws.append(names)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def save(wb, name):
    wb.save(name)
    print("создан", name)


def main():
    # справочник размеров
    wb = Workbook()
    ws = wb.active
    ws.title = "Sizes"
    header(ws, ["Размер"])
    for s in SIZES:
        ws.append([s])
    save(wb, "Sizes_import.xlsx")

    # пользователи
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    header(ws, ["ФИО", "Логин", "Роль"])
    for row in USERS:
        ws.append(list(row))
    save(wb, "Users_import.xlsx")

    # товары (картинки есть не у всех, у последних шести файла нет, там будет заглушка)
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    header(ws, ["Наименование", "Категория", "Производитель", "Описание", "Состав", "Цена", "Изображение"])
    for i, (name, cat, man, desc, comp, price) in enumerate(PRODUCTS, 1):
        ws.append([name, cat, man, desc, comp, price, "IMG_%03d.png" % i])
    save(wb, "Products_import.xlsx")

    # остатки: по 2-4 размера на модель, у нескольких моделей совсем мало
    stock = []
    scarce = {3, 8, 14, 22, 29}          # номера моделей, у которых остатки маленькие
    for i, (name, cat, man, *_rest) in enumerate(PRODUCTS, 1):
        lo, hi = SIZE_RANGE[cat]
        count = 2 if i in scarce else random.choice([3, 3, 3, 4])
        sizes = sorted(random.sample(range(lo, hi + 1), count))
        for s in sizes:
            qty = random.choice([1, 1, 2]) if i in scarce else random.choice([2, 3, 4, 5, 6, 8, 10, 12, 15])
            stock.append([name, man, s, qty])
    # две строки с опечаткой в названии: такого товара нет, импорт должен их пропустить и написать об этом
    stock.append(["Туфли-лодочки Clasic", "Salamander", 38, 4])
    stock.append(["Кроссовки Air Max 720", "Nike", 42, 3])
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    header(ws, ["Наименование", "Производитель", "Размер", "Количество"])
    for row in stock:
        ws.append(row)
    save(wb, "Stock_Items_import.xlsx")

    # заказы: 10 штук, 30 позиций; 5 заказов в прошлом месяце, остальные раньше
    today = date.today()
    first_this = today.replace(day=1)
    last_month_end = first_this - timedelta(days=1)
    lm = last_month_end.replace(day=1)
    dates = [lm.replace(day=d) for d in (3, 8, 12, 17, 24)]
    older = [(lm - timedelta(days=40)).replace(day=d) for d in (5, 14, 20)] + [(lm - timedelta(days=75)).replace(day=d) for d in (9, 22)]
    all_dates = sorted(dates + older, reverse=True)
    clients = [u[0] for u in USERS if u[2] == "Клиент"] + GUESTS
    price_of = {(p[0], p[2]): p[5] for p in PRODUCTS}
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    header(ws, ["Номер заказа", "Дата заказа", "ФИО клиента", "Наименование", "Производитель", "Размер", "Количество", "Цена"])
    counts = [3, 3, 4, 2, 3, 3, 3, 3, 3, 3]          # 30 позиций
    pool = [r for r in stock if r[0] in {p[0] for p in PRODUCTS}]
    for n, (d, cnt) in enumerate(zip(all_dates, counts), 1):
        client = clients[(n * 3) % len(clients)] if n > 2 else GUESTS[n - 1]
        picked = random.sample(pool, cnt)
        for name, man, size, _q in picked:
            price = price_of[(name, man)]
            if d < lm:                                   # раньше цены были чуть ниже
                price = int(round(price * 0.95, -1))
            ws.append([1000 + n, datetime(d.year, d.month, d.day), client, name, man, size, random.choice([1, 1, 1, 2]), price])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "DD.MM.YYYY"
    save(wb, "Orders_import.xlsx")
    print("строк остатков:", len(stock), "| товаров:", len(PRODUCTS))


if __name__ == "__main__":
    main()
